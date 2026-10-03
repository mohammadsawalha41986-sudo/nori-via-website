"""Build compact, bilingual library previews carrying the supplied original logo."""
from pathlib import Path
import json, re
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import arabic_reshaper
from bidi.algorithm import get_display
ROOT=Path(__file__).resolve().parent.parent
LOGO=ROOT/'public/brand/noriva-original.jpg'
FONT='/usr/share/fonts/truetype/dejavu/'
pdfmetrics.registerFont(TTFont('Brief',FONT+'DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('BriefBold',FONT+'DejaVuSans-Bold.ttf'))
def ar(s): return get_display(arabic_reshaper.reshape(s))
def ar_wrap(text,width=480,size=10):
 lines=[];line=''
 for w in text.split():
  trial=(line+' '+w).strip()
  if line and pdfmetrics.stringWidth(ar(trial),'Brief',size)>width: lines.append(ar(line));line=w
  else: line=trial
 if line: lines.append(ar(line))
 return '<br/>'.join(escape(s) for s in lines)
def para(d,text,arabic=False,style=None):
 p=d.add_paragraph(text,style)
 if arabic:
  p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
  bidi=OxmlElement('w:bidi');p._p.get_or_add_pPr().append(bidi)
 return p
CTA_EN='For a complete project study and implementation plan: norivaglobal.com/ar/start-a-project'
CTA_AR='لإكمال دراسة مشروعك وخطة التنفيذ، تواصل مع نوريفا عبر الموقع.'
def brief_doc(doc):
 d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.7)
 d.styles['Title'].font.color.rgb=RGBColor(0,0,0)
 for st in d.styles:
  pp=st.element.find(qn('w:pPr'))
  if pp is not None:
   for border in list(pp.findall(qn('w:pBdr'))): pp.remove(border)
 normal=d.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(5)
 d.add_picture(str(LOGO),width=Inches(2.4))
 para(d,re.sub(r'[^\w\s]',' ',doc['titleEn']),style='Title')
 para(d,doc['titleAr'],True)
 para(d,'A brief starting worksheet. Record your project evidence, then ask NORIVA to complete the review.')
 para(d,'نموذج بداية مختصر. سجّل بيانات مشروعك ثم شاركها مع نوريفا لإكمال المراجعة.',True)
 t=d.add_table(rows=2,cols=2);t.style='Table Grid'
 for r,(a,b) in zip(t.rows,[('Project / المشروع','Date / التاريخ'),('________________','________________')]):r.cells[0].text=a;r.cells[1].text=b
 for en,arabic in doc['includes'][:3]:
  para(d,re.sub(r'[^\w\s]',' ',en),style='Heading 2');para(d,arabic,True)
  para(d,'Evidence / الدليل: ______________________________________________')
  para(d,'Decision or action / القرار أو الإجراء: ______________________________')
 para(d,'Next step',style='Heading 2');para(d,'Owner / المسؤول: __________  Review date / تاريخ المراجعة: __________')
 para(d,CTA_AR,True);para(d,CTA_EN)
 s.footer.paragraphs[0].text='NORIVA  |  Brief working sample  |  norivaglobal.com'
 d.save(ROOT/'resources/library/docx'/doc['file'])
def brief_pdf(doc,target=None,rows=None,decision=None):
 base=ParagraphStyle('Base',fontName='Brief',fontSize=10,leading=15,textColor=colors.HexColor('#15223b'),spaceAfter=6)
 rt=ParagraphStyle('Arabic',parent=base,alignment=2)
 title=ParagraphStyle('Title',parent=base,fontName='BriefBold',fontSize=18,leading=23,spaceAfter=10)
 heading=ParagraphStyle('Heading',parent=base,fontName='BriefBold',fontSize=11,spaceBefore=10)
 P=lambda s:Paragraph(escape(s),base)
 A=lambda s:Paragraph(ar_wrap(s),rt)
 story=[Image(str(LOGO),width=175,height=58,hAlign='LEFT'),Spacer(1,16),Paragraph(escape(doc['titleEn']),title),A(doc['titleAr']),
 P('Brief working sample. Use verified project data to complete the review.'),A('نموذج عمل مختصر. استخدم بيانات موثقة لمشروعك لإكمال المراجعة.'),Spacer(1,10)]
 if rows:
  story += [P('Illustrative figures in SAR, excluding VAT. These are fictional inputs.'),A('أرقام افتراضية بالريال السعودي دون الضريبة؛ لا تمثل نتائج عميل.')]
  table=Table([[P('Item'),A('البند'),P('Value')]]+[[P(en),A(a),P(value)] for en,a,value in rows],colWidths=[190,205,90])
 else:
  story += [P('Project: _______________________   Review date: _______________________')]
  table=Table([[P('Review point'),A('نقطة المراجعة'),A('الدليل والإجراء')]]+[[P(en),A(a),'________________'] for en,a in doc['includes'][:3]],colWidths=[190,195,100])
 table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#f2f3f6')),('GRID',(0,0),(-1,-1),.5,colors.HexColor('#d5d9e0')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),10)]));story += [table,Spacer(1,16)]
 if decision:
  story.extend([Paragraph('Decision to test',heading),P(decision[0]),A(decision[1])])
 else:
  story.extend([Paragraph('Action and follow up',heading),P('Owner: _____________________   Review date: _____________________'),A('النتيجة أو السؤال المطلوب استكماله: __________________________________')])
 story.extend([Spacer(1,18),A(CTA_AR),P(CTA_EN)])
 path=target or ROOT/'resources/library/pdf'/doc['file']
 SimpleDocTemplate(str(path),pagesize=A4,leftMargin=40,rightMargin=40,topMargin=32,bottomMargin=32,title=doc['titleEn'],author='NORIVA').build(story)
manifest=json.load(open(ROOT/'scripts/brief-resource-manifest.json'))['documents']
for doc in manifest:
 if doc['kind']=='docx':brief_doc(doc)
 elif doc['kind']=='pdf':brief_pdf(doc)
samples=[
 ('sample-delivery-contribution-review','Delivery contribution review','مراجعة مساهمة التوصيل',[
 ('Order revenue','إيراد الطلب','60.00'),('Food cost','تكلفة الطعام','18.00'),('Platform commission','عمولة المنصة','15.00'),('Packaging','التغليف','3.00'),('Restaurant discount','خصم يتحمله المطعم','6.00'),('Other variable fulfilment','تكلفة تنفيذ متغيرة أخرى','1.50'),('Order contribution','مساهمة الطلب','16.50 / 27.5%')],('Contribution covers fixed costs; it is not net profit. Verify the commission base, test a bundle, then compare contribution at the expected order volume.','المساهمة تغطي التكاليف الثابتة وليست صافي الربح. تحقق من أساس العمولة واختبر باقة ثم قارن المساهمة عند حجم الطلبات المتوقع.')),
 ('sample-monthly-restaurant-review','Monthly restaurant review','المراجعة الشهرية للمطعم',[
 ('Monthly revenue','الإيراد الشهري','300,000'),('Food and beverage cost','تكلفة الطعام والمشروبات','90,000'),('Labour','العمالة','75,000'),('Occupancy','الإشغال','30,000'),('Other operating costs','تكاليف تشغيل أخرى','60,000'),('Operating result','النتيجة التشغيلية','45,000 / 15%')],('Result before depreciation, interest and tax. Reconcile recipe yield, match labour hours to demand and verify month-end accruals. Example percentages are not targets.','النتيجة قبل الإهلاك والفوائد والضرائب. طابق مردود الوصفة وساعات العمل مع الطلب وتحقق من الاستحقاقات. النسب توضيحية وليست أهدافًا.')),
 ('sample-menu-decision-review','Menu decision review','مراجعة قرارات القائمة',[
 ('Chicken: price / cost / sold','الدجاج: السعر والتكلفة والكمية','48 / 15 / 240'),('Pasta: price / cost / sold','الباستا: السعر والتكلفة والكمية','42 / 13 / 180'),('Vegetables: price / cost / sold','الخضروات: السعر والتكلفة والكمية','38 / 11 / 70'),('Total portions sold','إجمالي الحصص المباعة','490'),('Ingredient-only contribution','المساهمة بعد المكونات فقط','15,030')],('Contribution excludes labour, waste and channel costs. Verify recipe yield and test visibility before removing a low-demand dish. Change one factor at a time.','المساهمة لا تخصم العمالة والهدر وتكلفة القنوات. تحقق من مردود الوصفة واختبر ظهور الصنف قبل حذفه. غيّر عاملًا واحدًا في كل مرة.'))]
for slug,en,a,rows,decision in samples:
 brief_pdf({'titleEn':en,'titleAr':a},ROOT/'resources/library/pdf'/f'{slug}.pdf',rows,decision)
print('Built 19 compact DOCX, 20 compact PDF guides and 3 distinct worked examples')
